"""deck_merge — PowerPoint での手編集とスクリプト生成の共同編集（3-way マージ）＋時刻付きバックアップ。

考え方（Git の 3-way マージと同じ）:
    base   = 前回スクリプトが生成した内容（output/bg/<name>.base.pptx。要約はマージのたびに今のコードで作り直す
             ので、要約の方式を変えても誤った「衝突」が出ない。<name>.base.json には sha・時刻だけを残す）
    theirs = あなたが PowerPoint で編集・保存した <name>.pptx
    mine   = 今回スクリプトが生成した内容（output/bg/<name>.gen.pptx）
部品 = スライド（キー = タイトル#出現回数 を cSld@name に保存）× 図形（名前 g:pic1 / g:txt2 / TITLE / MSG / REF / QUOTE …）。

規則（部品ごと）:
    あなたが触っていない（theirs == base）     → mine に差し替える
    あなたが中身を変えた（theirs != base）     → あなたの版を残す。mine も変わっていれば「衝突」として報告
    位置・大きさ                              → あなたが動かしていなければ mine に、動かしていればあなたの版
    あなたが足した図形・スライド（base に無い）→ 触らない
    あなたが消した図形・スライド              → 戻さない（mine が変わっていれば報告）
    スクリプトが消した部品                    → あなたが触っていなければ消す、触っていれば残して報告
    スライドの順番                            → あなたが並べ替えていなければ生成側の順番、並べ替えていればあなたの順番
変更の判定は XML そのものでなく意味の要約（文字・書式・画像ハッシュ・位置）で行う
（PowerPoint は保存時に lang・dirty などの属性を足すので、XML 比較だと誤って「衝突」になる）。

保存（slide_template.save → save_coedit）:
    1. PowerPoint が開いている（~$<name>.pptx がある）なら書き込まずに止める
    2. 既存の <name>.pptx を output/bg/<name>_YYYYMMDD-HHMMSS.pptx へ必ずコピー
    3. 既存ファイルが前回の生成のまま → mine で置き換え。手編集あり → マージして上書き
    4. 今回の生成を output/bg/<name>.base.pptx（次回の基準）にし、base.json（sha・時刻）を更新、
       マージ報告を output/bg/<name>_<時刻>.merge.txt に残す
"""
from __future__ import annotations

import copy
import datetime as _dt
import hashlib
import io
import json
import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

BG_DIRNAME = "bg"
KEEP_NAMES = ("TITLE", "MSG", "REF", "QUOTE")
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ─── 識別子 ────────────────────────────────────────────────────────────────
def _first_text(slide):
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip():
            return shp.text_frame.text.strip().split("\n")[0]
    return "(no text)"


def _slide_label(slide):
    for shp in slide.shapes:
        if shp.name == "TITLE" and shp.has_text_frame:
            return shp.text_frame.text.strip().split("\n")[0]
    return _first_text(slide)


def _kind(shp):
    tag = etree.QName(shp._element).localname
    if tag == "pic":
        return "pic"
    if tag == "graphicFrame":
        return "tbl" if getattr(shp, "has_table", False) and shp.has_table else "gfx"
    if tag == "cxnSp":
        return "cxn"
    if tag == "grpSp":
        return "grp"
    return "txt" if shp.has_text_frame and shp.text_frame.text.strip() else "shp"


def slide_key(slide):
    """保存済みのキー（cSld@name）。無ければ None。"""
    return slide._element.cSld.get("name") or None


def tag_deck(prs):
    """生成したデッキに、スライドのキー（タイトル#出現回数）と図形名（g:<種類><番号>）を付ける。"""
    seen = {}
    for slide in prs.slides:
        lab = _slide_label(slide)[:60]
        seen[lab] = seen.get(lab, 0) + 1
        slide._element.cSld.set("name", f"{lab}#{seen[lab]}")
        cnt = {}
        for shp in slide.shapes:
            if shp.name in KEEP_NAMES:
                continue
            k = _kind(shp)
            cnt[k] = cnt.get(k, 0) + 1
            shp.name = f"g:{k}{cnt[k]}"


# ─── 意味の要約（PowerPoint の保存で変わらない量だけ） ────────────────────────────
def _fill_sig(parent):
    """solidFill の色（srgb・テーマ色とその加工）を文字列に。"""
    if parent is None:
        return None
    f = parent.find(f"{_A}solidFill")
    if f is None or len(f) == 0:
        return None
    c = f[0]
    mods = ",".join(f"{etree.QName(m).localname}={m.get('val')}" for m in c)
    return f"{etree.QName(c).localname}:{(c.get('val') or '').upper()}" + (f"[{mods}]" if mods else "")


def _run_sig(el):
    rpr = el.find(f"{_A}rPr")
    a = {}
    if rpr is not None:
        for k in ("sz", "b", "i", "u", "baseline", "strike", "cap"):
            if rpr.get(k) is not None:
                a[k] = rpr.get(k)
        col = _fill_sig(rpr)
        if col:
            a["c"] = col
        for tag in ("latin", "ea"):
            f = rpr.find(f"{_A}{tag}")
            if f is not None and f.get("typeface"):
                a[tag] = f.get("typeface")
        if rpr.find(f"{_A}highlight") is not None:
            a["hl"] = 1
    t = el.find(f"{_A}t")
    return [t.text if t is not None else "", a]


def _ppr_sig(ppr):
    """段落の書式: 配置・レベル・字下げ・行間・段落前後・箇条書き。"""
    if ppr is None:
        return None
    d = {k: ppr.get(k) for k in ("algn", "lvl", "marL", "indent") if ppr.get(k) is not None}
    for tag in ("lnSpc", "spcBef", "spcAft"):
        e = ppr.find(f"{_A}{tag}")
        if e is not None and len(e):
            d[tag] = f"{etree.QName(e[0]).localname}={e[0].get('val')}"
    for tag in ("buChar", "buAutoNum", "buNone", "buFont", "buSzPct"):
        e = ppr.find(f"{_A}{tag}")
        if e is not None:
            d[tag] = e.get("char") or e.get("type") or e.get("typeface") or e.get("val") or 1
    return d or None


def _txbody_sig(txbody):
    """段落ごとの [配置, [文字, 書式], …]。同じ書式の隣り合う run は結合する
    （PowerPoint は保存時に run を結合・分割するので、run の切れ目は比べない）。空の run は捨てる。"""
    out = []
    for p in txbody.iter(f"{_A}p"):
        para = [_ppr_sig(p.find(f"{_A}pPr"))]
        for ch in p:
            tag = etree.QName(ch).localname
            if tag in ("r", "fld"):
                seg = _run_sig(ch)
            elif tag == "br":
                seg = ["\n", {}]
            else:
                continue
            if not seg[0]:
                continue
            if len(para) > 1 and para[-1][1] == seg[1]:
                para[-1] = [para[-1][0] + seg[0], seg[1]]
            else:
                para.append(seg)
        out.append(para)
    return out


def _content_sig(shp):
    el = shp._element
    k = _kind(shp)
    if k == "pic":
        try:
            return "img:" + hashlib.sha256(shp.image.blob).hexdigest()
        except Exception:      # noqa: BLE001  リンク画像など
            return "img:?"
    if k == "tbl":
        tbl = el.xpath(".//a:tbl")[0]
        grid = [g.get("w") for g in tbl.iter(f"{_A}gridCol")]
        rows = []
        for tr in tbl.iter(f"{_A}tr"):
            cells = []
            for tc in tr.iter(f"{_A}tc"):
                span = {k: tc.get(k) for k in ("gridSpan", "rowSpan", "hMerge", "vMerge") if tc.get(k)}
                tcpr = tc.find(f"{_A}tcPr")
                cells.append([_txbody_sig(tc.find(f"{_A}txBody")), span, _fill_sig(tcpr)])
            rows.append([tr.get("h"), cells])
        return "tbl:" + hashlib.sha256(json.dumps([grid, rows], ensure_ascii=False).encode()).hexdigest()
    tx = el.find(f".//{{http://schemas.openxmlformats.org/presentationml/2006/main}}txBody")
    sig = _txbody_sig(tx) if tx is not None else []
    return "txt:" + hashlib.sha256(json.dumps(sig, ensure_ascii=False).encode()).hexdigest()


def _style_sig(shp):
    """図形の書式（枠線・塗り・テーマ書式）の要約。位置（a:xfrm）は _geom で別に見るので除く。"""
    el = shp._element
    parts = []
    for tag in ("p:spPr", "p:style"):
        x = el.find(qn(tag))
        if x is None:
            parts.append("")
            continue
        x = copy.deepcopy(x)
        for xf in x.findall(qn("a:xfrm")):
            x.remove(xf)
        parts.append(etree.tostring(x, method="c14n").decode())
    return "sty:" + hashlib.sha256("|".join(parts).encode()).hexdigest()


def _replace_style(t_shp, m_shp):
    """書式（spPr の xfrm 以外・p:style）を生成側に差し替える。位置・大きさはあなたの版のまま。"""
    t_el, m_el = t_shp._element, m_shp._element
    t_sp, m_sp = t_el.find(qn("p:spPr")), m_el.find(qn("p:spPr"))
    if t_sp is not None and m_sp is not None:
        new = copy.deepcopy(m_sp)
        for xf in new.findall(qn("a:xfrm")):
            new.remove(xf)
        old_xf = t_sp.find(qn("a:xfrm"))
        if old_xf is not None:
            new.insert(0, copy.deepcopy(old_xf))
        t_sp.getparent().replace(t_sp, new)
    t_st, m_st = t_el.find(qn("p:style")), m_el.find(qn("p:style"))
    if t_st is not None:
        t_el.remove(t_st)
    if m_st is not None:
        t_el.append(copy.deepcopy(m_st))


def _plain_text(shp):
    if shp.has_text_frame:
        return shp.text_frame.text
    if getattr(shp, "has_table", False) and shp.has_table:
        return " | ".join(c.text for row in shp.table.rows for c in row.cells)
    return ""


def _geom(shp):
    try:
        return [int(shp.left), int(shp.top), int(shp.width), int(shp.height)]
    except TypeError:
        return None


def _notes(slide):
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def snapshot(prs):
    """デッキの要約 {slide_key: {"order", "layout", "notes", "shapes": {name: {kind, content, geom, text}}}}。"""
    snap = {}
    for i, slide in enumerate(prs.slides):
        key = slide_key(slide)
        if key is None or key in snap:
            continue
        shapes = {}
        for shp in slide.shapes:
            if shp.name in shapes:
                continue
            shapes[shp.name] = dict(kind=_kind(shp), content=_content_sig(shp), geom=_geom(shp),
                                    style=_style_sig(shp), text=_plain_text(shp)[:200])
        snap[key] = dict(order=i, layout=slide.slide_layout.name, notes=_notes(slide), shapes=shapes)
    return snap


# ─── 書き換えの部品 ─────────────────────────────────────────────────────────
def _next_id(slide):
    ids = [int(x) for x in slide._element.xpath(".//p:cNvPr/@id")]
    return max(ids, default=1) + 1


def _copy_shape(src_shp, dst_slide):
    """生成側の図形を、編集版のスライドへ複製する（画像は画像パートごと）。"""
    k = _kind(src_shp)
    if k == "pic":
        new = dst_slide.shapes.add_picture(io.BytesIO(src_shp.image.blob), src_shp.left, src_shp.top,
                                           src_shp.width, src_shp.height)
        new.name = src_shp.name
        return new
    el = copy.deepcopy(src_shp._element)
    nid = _next_id(dst_slide)
    for c in el.xpath(".//p:cNvPr"):        # グループ内の子にも 1 つずつ別の id を振る
        c.set("id", str(nid))
        nid += 1
    dst_slide.shapes._spTree.append(el)
    return el


def _replace_picture(t_shp, m_shp, report, where):
    """画像だけ差し替える（位置・大きさはあなたの版）。縦横比が変われば幅を保って高さを合わせる。"""
    slide_part = t_shp.part
    _, rid = slide_part.get_or_add_image_part(io.BytesIO(m_shp.image.blob))
    blip = t_shp._element.xpath(".//a:blip")[0]
    blip.set(qn("r:embed"), rid)
    src = t_shp._element.xpath(".//a:srcRect")
    if src and any(src[0].get(k) for k in ("l", "t", "r", "b")):
        report.append(f"  [注意] {where}: あなたが図を切り出し（トリミング）しているので、新しい図で切り出し範囲を確認してください")
    ar_new = m_shp.width / m_shp.height
    ar_t = t_shp.width / t_shp.height
    if abs(ar_new / ar_t - 1) > 0.02:
        t_shp.height = int(t_shp.width / ar_new)
        report.append(f"  [note] {where}: 図の縦横比が変わったので、幅を保って高さを合わせた")


def _replace_text(t_shp, m_shp):
    t_tx = t_shp._element.find(qn("p:txBody"))
    m_tx = m_shp._element.find(qn("p:txBody"))
    if t_tx is not None and m_tx is not None:
        t_tx.getparent().replace(t_tx, copy.deepcopy(m_tx))


def _replace_table(t_shp, m_shp):
    t_tbl = t_shp._element.xpath(".//a:tbl")[0]
    m_tbl = m_shp._element.xpath(".//a:tbl")[0]
    t_tbl.getparent().replace(t_tbl, copy.deepcopy(m_tbl))


def _set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _shape_map(slide):
    m, dup = {}, []
    for shp in slide.shapes:
        if shp.name in m:
            dup.append(shp.name)
            continue
        m[shp.name] = shp
    return m, dup


# ─── マージ本体 ─────────────────────────────────────────────────────────────
def merge(theirs_path, mine_path, base):
    """theirs（あなたの編集版）に mine（今回の生成）を 3-way で反映した Presentation と報告行を返す。"""
    th = Presentation(str(theirs_path))
    mi = Presentation(str(mine_path))
    bsnap = base["slides"]
    rep, n_upd, n_conf, n_add = [], 0, 0, 0
    th_by_key = {}
    for s in th.slides:
        k = slide_key(s)
        if k and k not in th_by_key:
            th_by_key[k] = s
        elif k:
            rep.append(f"[注意] スライド「{k}」が複数ある（複製?）。最初の 1 枚だけを対象にした")
    mine_keys = [slide_key(s) for s in mi.slides]
    prev_th = None                                        # 新しいスライドを差し込む位置
    for ms in mi.slides:
        key = slide_key(ms)
        b = bsnap.get(key)
        ts = th_by_key.get(key)
        if ts is None:
            if b is None:                                 # 生成側で増えたスライド → 追加
                lay = next((l for l in th.slide_layouts if l.name == ms.slide_layout.name), th.slide_layouts[-1])
                ts = add_slide_safe(th, lay)
                for ph in list(ts.placeholders):
                    ph._element.getparent().remove(ph._element)
                ts._element.cSld.set("name", key)
                for shp in ms.shapes:
                    _copy_shape(shp, ts)
                if ms.has_notes_slide:
                    _set_notes(ts, _notes(ms))
                _move_after(th, ts, _skip_user_followers(th, prev_th, bsnap, mine_keys))
                rep.append(f"[追加] スライド「{key}」")
                n_add += 1
                prev_th = ts
            else:
                rep.append(f"[そのまま] スライド「{key}」はあなたが削除したので戻さない")
            continue
        prev_th = ts
        tmap, dup = _shape_map(ts)
        for d in dup:
            rep.append(f"[注意] 「{key}」の図形名 {d} が複数ある（複製?）。最初の 1 つだけを対象にした")
        bsh = b["shapes"] if b else {}
        mmap, _ = _shape_map(ms)
        for name, mshp in mmap.items():
            where = f"「{key}」/{name}"
            msig = dict(content=_content_sig(mshp), geom=_geom(mshp))
            bs = bsh.get(name)
            tshp = tmap.get(name)
            if tshp is None:
                if bs is None:
                    _copy_shape(mshp, ts)
                    rep.append(f"[追加] {where}")
                    n_add += 1
                elif msig["content"] != bs["content"]:
                    rep.append(f"[そのまま] {where} はあなたが削除済み（生成側は内容を変更: {_plain_text(mshp)[:40]!r}）")
                continue
            if bs is None:
                rep.append(f"[衝突] {where} は前回の生成に無い名前。あなたの版を残した")
                n_conf += 1
                continue
            tc, tg = _content_sig(tshp), _geom(tshp)
            if msig["content"] != bs["content"]:
                if tc == bs["content"]:
                    k = _kind(mshp)
                    done = True
                    if k == "pic" and _kind(tshp) == "pic":
                        _replace_picture(tshp, mshp, rep, where)
                    elif k == "tbl" and _kind(tshp) == "tbl":
                        _replace_table(tshp, mshp)
                    elif mshp.has_text_frame and tshp.has_text_frame:
                        _replace_text(tshp, mshp)
                    else:
                        done = False
                    if done:
                        rep.append(f"[更新] {where}" + (f": {_plain_text(mshp)[:50]!r}" if k != "pic" else "（図）"))
                        n_upd += 1
                    else:
                        rep.append(f"[衝突] {where}: 種類（{k}）を自動で差し替えられない。手で確認してください")
                        n_conf += 1
                else:
                    rep.append(f"[衝突] {where}: あなたの版を残した。生成側の案 = {_plain_text(mshp)[:80]!r}")
                    n_conf += 1
            ms_sty = _style_sig(mshp)                      # 書式（枠線の色・太さ・塗り）も 3-way で反映（10/29）
            if "style" in bs and ms_sty != bs["style"] and _kind(mshp) != "pic":
                if _style_sig(tshp) == bs["style"]:
                    _replace_style(tshp, mshp)
                    rep.append(f"[更新] {where} の書式（枠線・塗り）")
                    n_upd += 1
                else:
                    rep.append(f"[そのまま] {where} はあなたが書式を変えたので、書式はあなたの版")
            if msig["geom"] != bs["geom"] and msig["geom"] is not None:
                if tg == bs["geom"]:
                    tshp.left, tshp.top, tshp.width, tshp.height = msig["geom"]
                    rep.append(f"[更新] {where} の位置・大きさ")
                    n_upd += 1
                else:
                    rep.append(f"[そのまま] {where} はあなたが動かしたので位置・大きさはあなたの版")
        for name, bs in bsh.items():                      # 生成側で消えた図形
            if name in mmap or name not in tmap:
                continue
            tshp = tmap[name]
            if _content_sig(tshp) == bs["content"] and _geom(tshp) == bs["geom"]:
                tshp._element.getparent().remove(tshp._element)
                rep.append(f"[削除] 「{key}」/{name}（生成側で不要になった）")
            else:
                rep.append(f"[衝突] 「{key}」/{name} は生成側で不要になったが、あなたが編集したので残した")
                n_conf += 1
        mn = _notes(ms)
        if b is not None and mn != b["notes"]:
            if _notes(ts) == b["notes"]:
                _set_notes(ts, mn)
                rep.append(f"[更新] 「{key}」のノート")
                n_upd += 1
            else:
                rep.append(f"[衝突] 「{key}」のノート: あなたの版を残した")
                n_conf += 1
    for key, ts in th_by_key.items():                     # 生成側で消えたスライド
        if key in mine_keys or key not in bsnap:
            continue
        if snapshot_slide_equal(ts, bsnap[key]):
            _delete_slide(th, ts)
            rep.append(f"[削除] スライド「{key}」（生成側で不要になった）")
        else:
            rep.append(f"[衝突] スライド「{key}」は生成側で不要になったが、あなたが編集したので残した")
            n_conf += 1
    _reorder(th, mine_keys, bsnap, rep)
    user_slides = [slide_key(s) or _user_label(s) for s in th.slides if (slide_key(s) or "") not in bsnap
                   and (slide_key(s) or "") not in mine_keys]
    if user_slides:
        rep.append(f"[そのまま] あなたが足したスライド {len(user_slides)} 枚: {', '.join(x[:30] for x in user_slides)}")
    head = f"更新 {n_upd}・追加 {n_add}・衝突 {n_conf}"
    return th, [head] + rep


def _skip_user_followers(th, slide, bsnap, mine_keys):
    """slide の直後に続く「あなたが足したスライド」の最後を返す（新しいスライドをその後ろに入れ、あなたのスライドを
    直前のスライドから引き離さないため）。slide が None なら None。"""
    if slide is None:
        return None
    sl = list(th.slides)
    i = next(j for j, x in enumerate(sl) if x._element is slide._element)
    while i + 1 < len(sl) and (slide_key(sl[i + 1]) or "") not in bsnap and (slide_key(sl[i + 1]) or "") not in mine_keys:
        i += 1
    return sl[i]


def _reorder(th, mine_keys, bsnap, rep):
    """あなたが生成スライドを並べ替えていなければ、生成側の順番に合わせる。
    あなたが足したスライドは、直前にあったスライドの後ろに付いたまま一緒に動く。並べ替えていれば、あなたの順番を残す。"""
    lst = th.slides._sldIdLst
    els = list(lst)
    keys = [slide_key(s) for s in th.slides]
    common = [k for k in keys if k in bsnap and k in mine_keys]
    if common != sorted(common, key=lambda k: bsnap[k]["order"]):
        rep.append("[そのまま] あなたがスライドを並べ替えているので、順番はあなたの版")
        return
    head, groups, cur = [], {}, None
    for el, k in zip(els, keys):
        if k in mine_keys and k not in groups:
            cur = k
            groups[k] = [el]
        elif cur is None:
            head.append(el)
        else:
            groups[cur].append(el)
    new = head + [e for k in mine_keys if k in groups for e in groups[k]]
    if [id(e) for e in new] != [id(e) for e in els]:
        for e in els:
            lst.remove(e)
        for e in new:
            lst.append(e)
        rep.append("[並べ替え] 生成側の順番に合わせた（あなたが足したスライドは直前のスライドと一緒に移動）")


def add_slide_safe(prs, layout):
    """python-pptx の add_slide は新しいスライドを slide<枚数+1>.xml と名付けるので、削除の後だと既存と衝突して
    ファイルが壊れる。空いている番号を使って追加する。"""
    from pptx.parts.slide import SlidePart
    pkg = prs.part.package
    part = SlidePart.new(pkg.next_partname("/ppt/slides/slide%d.xml"), pkg, layout.part)
    rid = prs.part.relate_to(part, _SLIDE_RT)
    slide = part.slide
    slide.shapes.clone_layout_placeholders(layout)
    prs.slides._sldIdLst.add_sldId(rid)
    return slide


def _user_label(slide):
    texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    texts = [t for t in texts if not t.isdigit()] or texts or ["(no text)"]
    return max(texts, key=len).split("\n")[0]


def snapshot_slide_equal(slide, bs):
    for shp in slide.shapes:
        b = bs["shapes"].get(shp.name)
        if b is None or b["content"] != _content_sig(shp) or b["geom"] != _geom(shp):
            return False
    return len(list(slide.shapes)) == len(bs["shapes"]) and _notes(slide) == bs["notes"]


def _move_after(prs, slide, after):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    me = next(x for x in ids if x.get(qn("r:id")) == prs.part.relate_to(slide.part, _SLIDE_RT))
    lst.remove(me)
    if after is None:
        lst.insert(0, me)
        return
    aid = prs.part.relate_to(after.part, _SLIDE_RT)
    pos = next(i for i, x in enumerate(lst) if x.get(qn("r:id")) == aid)
    lst.insert(pos + 1, me)


def _delete_slide(prs, slide):
    rid = prs.part.relate_to(slide.part, _SLIDE_RT)
    lst = prs.slides._sldIdLst
    for x in list(lst):
        if x.get(qn("r:id")) == rid:
            lst.remove(x)
    prs.part.drop_rel(rid)


_SLIDE_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"


# ─── 保存の入口 ─────────────────────────────────────────────────────────────
def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def lock_path(out):
    out = Path(out)
    return out.parent / ("~$" + out.name)


def backup(out, stamp=None):
    """<dir>/bg/<stem>_<YYYYMMDD-HHMMSS><suffix> にコピーしてパスを返す（同じ秒なら -2, -3…）。"""
    out = Path(out)
    bg = out.parent / BG_DIRNAME
    bg.mkdir(exist_ok=True)
    stamp = stamp or _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    dst, k = bg / f"{out.stem}_{stamp}{out.suffix}", 2
    while dst.exists():
        dst, k = bg / f"{out.stem}_{stamp}-{k}{out.suffix}", k + 1
    shutil.copy2(out, dst)
    return dst


def _atomic_save(prs_or_path, dst):
    """同じフォルダの一時ファイルに書いてから os.replace（途中で止まっても dst は壊れない）。"""
    import os
    dst = Path(dst)
    tmp = dst.parent / (dst.stem + ".__tmp__" + dst.suffix)
    if isinstance(prs_or_path, (str, Path)):
        shutil.copy2(prs_or_path, tmp)
    else:
        prs_or_path.save(str(tmp))
    os.replace(tmp, dst)


def save_coedit(prs, out_path, init=False):
    """共同編集つき保存。戻り値 = dict(mode, backup, report_path, report)。
    init=True: 基準ファイル（<name>.base.json）が無いときに、既存の <name>.pptx を生成版で置き換えて共同編集を始める
    （既存ファイルは bg/ にバックアップされる）。既定（False）では、基準が無いのに既存ファイルがあれば止まる
    （手編集を黙って上書きしないため）。"""
    out = Path(out_path).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    if lock_path(out).exists():
        raise RuntimeError(f"PowerPoint で開いています（{lock_path(out).name} がある）。閉じてから再実行してください: {out}")
    stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    bg = out.parent / BG_DIRNAME
    bg.mkdir(exist_ok=True)
    tag_deck(prs)
    gen = bg / f"{out.stem}.gen.pptx"
    prs.save(str(gen))
    base_pptx = bg / f"{out.stem}.base.pptx"
    base_path = out.with_suffix(".base.json")
    base = json.loads(base_path.read_text()) if base_path.exists() else None
    if base is not None and base_pptx.exists():
        base["slides"] = snapshot(Presentation(str(base_pptx)))   # 基準の要約は今のコードで作り直す
    if out.exists() and (base is None or "slides" not in base) and not init:
        raise RuntimeError(
            f"基準ファイル（{base_path.name} と bg/{base_pptx.name}）が無いので、{out.name} に手編集があっても区別できません。"
            f"上書きせずに止めました。今回の生成は {gen} にあります。手編集が無いと分かっていれば save(..., coedit_init=True) で"
            "置き換えて共同編集を始めてください（既存ファイルは bg/ にバックアップされます）。")
    res = dict(mode="new", backup=None, report_path=None, report=[])
    if out.exists():
        res["backup"] = backup(out, stamp)
        # 丸ごと置き換えてよいのは「前回は生成版をそのまま書いた（pure）＋その後だれも触っていない」ときだけ。
        # 前回がマージなら out にはあなたの編集が入っているので、触っていなくても必ずマージする（置き換えると消える）
        pure_untouched = base is not None and base.get("pure", False) and _sha(out) == base.get("file_sha")
        if base is not None and "slides" in base and not pure_untouched:
            merged, rep = merge(out, gen, base)
            _atomic_save(merged, out)
            rp = bg / f"{out.stem}_{stamp}.merge.txt"
            rp.write_text("\n".join(rep) + "\n")
            res.update(mode="merge", report_path=rp, report=rep)
        else:
            _atomic_save(gen, out)
            res["mode"] = "replace" if base is not None else "replace(init)"
    else:
        _atomic_save(gen, out)
    pure = res["mode"] != "merge"
    _atomic_save(gen, base_pptx)                                      # 次回の基準 = 今回の生成
    tmpj = base_path.parent / (base_path.name + ".__tmp__")
    tmpj.write_text(json.dumps(dict(generated_at=stamp, file_sha=_sha(out), pure=pure,
                                    base_pptx=str(base_pptx)), ensure_ascii=False, indent=1))
    import os
    os.replace(tmpj, base_path)
    return res
